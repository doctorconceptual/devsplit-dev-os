from pathlib import Path

import audit


PROJECT = Path(__file__).resolve().parents[2]


def test_load_checklist_preserves_all_rows_and_method_counts():
    rows = audit.load_checklist(PROJECT / "cro-audit-checklist.xlsx")
    assert len(rows) == 92
    assert rows[0]["#"] == 1
    assert rows[-1]["#"] == 92
    assert {row["Method"] for row in rows} == {"DOM", "API", "Vision", "Manual"}


def test_first_pending_store_is_read_from_input_workbook(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "sites.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sites"
    sheet.append(["#", "Store URL", "Store Name", "Niche", "Status"])
    sheet.append([1, "https://loandcointeriors.com/", "Lo & Co Interiors", "furniture", "Pending"])
    workbook.save(path)

    store = audit.first_pending_store(path)
    assert store["Store URL"] == "https://loandcointeriors.com/"
    assert store["Status"] == "Pending"


def test_dom_classifier_only_answers_grounded_checks():
    rows = [
        {"#": 4, "Observation": "No star rating / review count near the product title", "Method": "DOM"},
        {"#": 5, "Observation": "No reviews on the store at all", "Method": "DOM"},
    ]
    snapshot = {"url": "https://example.test/product", "elements": [
        {"tag": "h1", "text": "Product"},
        {"tag": "div", "classes": "reviews", "text": "4.9 12 reviews"},
    ]}
    results = audit.answer_dom_rows(rows, {"product": snapshot})
    assert results[4]["exists"] == "N"
    assert results[4]["confidence"] == "high"
    assert results[5]["exists"] == "N"
    assert results[5]["evidence"]


def test_review_rows_reference_both_viewports():
    row = {"#": 25, "Observation": "Value proposition unclear above the fold", "Method": "Vision"}
    result = audit.review_result(row, Path("output/store"))
    assert result["exists"] == "review"
    assert "desktop" in result["evidence"]
    assert "mobile" in result["evidence"]


def test_niche_specific_manual_rows_remain_review():
    row = {"#": 16, "Category": "Product Page", "Observation": "Out-of-stock variants not handled", "How to check": "", "Email-ready line": "", "Impact": "Medium", "Applies to": "outfit/fashion", "Method": "Manual"}
    output = audit.assemble([row], {"Niche": "furniture"}, {"desktop-home": {"url": "https://example.test/"}}, Path("output/store"), {})
    assert output[0]["Exists?"] == "review"


def test_update_input_handles_header_on_first_row(tmp_path):
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "sites.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sites"
    sheet.append(["#", "Store URL", "Store Name", "Niche", "Status", "Date Audited", "Output File / Link", "Notes"])
    sheet.append([1, "https://example.test/", "Example", "furniture", "Pending", None, None, None])
    workbook.save(path)

    audit.update_input(path, "https://example.test/", "output/results.xlsx")
    row = list(load_workbook(path, data_only=True)["Sites"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    assert row[4] == "Done"
    assert row[6] == "output/results.xlsx"
