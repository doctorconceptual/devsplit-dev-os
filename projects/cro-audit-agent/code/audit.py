"""Single-store CRO audit pipeline: rendered DOM + PageSpeed + human review rows."""
from __future__ import annotations

import argparse
import json
import os
import re
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook

import browser
from pagespeed_api import evaluate as evaluate_pagespeed

METHODS = {"DOM", "API", "Vision", "Manual"}


def load_checklist(path: Path) -> list[dict]:
    ws = load_workbook(path, data_only=True)["Master Checklist"]
    values = list(ws.values)
    header_index = next(index for index, row in enumerate(values) if row and row[0] == "#")
    headers = list(values[header_index])
    rows = []
    for values_row in values[header_index + 1:]:
        if values_row[0] is None:
            continue
        row = dict(zip(headers, values_row))
        row["#"] = int(row["#"])
        if row["Method"] not in METHODS:
            raise ValueError(f"Unknown checklist method at #{row['#']}: {row['Method']}")
        rows.append(row)
    if len(rows) != 92 or [row["#"] for row in rows] != list(range(1, 93)):
        raise ValueError("Checklist must contain exactly 92 ordered rows")
    return rows


def first_pending_store(path: Path) -> dict:
    ws = load_workbook(path, data_only=True)["Sites"]
    header_row = next(row for row in ws.iter_rows() if row[1].value == "Store URL")
    headers = [cell.value for cell in header_row]
    for values in ws.iter_rows(min_row=header_row[0].row + 1, values_only=True):
        row = dict(zip(headers, values))
        if row.get("Status") == "Pending":
            return row
    raise ValueError("No Pending store found")


def slugify(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def _text(snapshots: dict) -> str:
    chunks = []
    for snapshot in snapshots.values():
        for element in snapshot.get("elements", []):
            chunks.extend(str(element.get(key) or "") for key in ("text", "classes", "label", "outer", "context"))
    return " ".join(chunks).lower()


def _elements(snapshots: dict) -> list[dict]:
    return [element for snapshot in snapshots.values() for element in snapshot.get("elements", [])]


def _has(snapshots: dict, pattern: str) -> bool:
    return bool(re.search(pattern, _text(snapshots), re.I))


def _usable(snapshots: dict) -> bool:
    return bool(snapshots) and any(snapshot.get("url") for snapshot in snapshots.values())


def _dom_presence(row: dict, snapshots: dict) -> tuple[bool | None, str]:
    observation = row["Observation"].lower()
    if not _usable(snapshots):
        return None, "Rendered evidence was unavailable; no DOM conclusion made."
    text = _text(snapshots)
    terms = {
        "star rating": r"\b(star|rating|reviews?|\d(?:\.\d)?\s*/\s*5)\b",
        "reviews": r"\b(review|reviews|testimonials?|customer stories)\b",
        "image": r"\b(img|image|gallery|photo|swatch)\b",
        "video": r"\b(video|youtube|vimeo|play)\b",
        "returns": r"\b(return|refund|guarantee|warranty)\b",
        "shipping": r"\b(shipping|delivery|deliver|dispatch|free shipping)\b",
        "urgency": r"\b(low stock|selling fast|limited|last one|only \d+)\b",
        "payment": r"\b(klarna|afterpay|shop pay|paypal|apple pay|google pay|visa|mastercard)\b",
        "cross-sell": r"\b(complete the look|you may also like|related|pairs well|recommended|frequently bought)\b",
        "faq": r"\b(faq|frequently asked|questions)\b",
        "wishlist": r"\b(wishlist|save for later|favourite|favorite)\b",
        "announcement": r"\b(free shipping|announcement|sale|offer)\b",
        "category": r"\b(shop|collections?|categories|furniture|hardware)\b",
        "filter": r"\b(filter|refine|sort by)\b",
        "sort": r"\b(sort by|best selling|price|newest)\b",
        "search": r"\b(search|search products)\b",
        "breadcrumb": r"\b(breadcrumb|home\s*[>/]|\s*/\s*.+\s*/\s*)\b",
        "contact": r"\b(contact|email us|live chat|chat)\b",
        "social": r"\b(instagram|facebook|pinterest|tiktok|social)\b",
        "email": r"\b(sign up|subscribe|newsletter|email|sms|discount)\b",
        "first-order": r"\b(first order|welcome|discount|save \d+%|offer)\b",
        "offer": r"\b(first order|welcome|discount|save \d+%|buy more|bundle)\b",
        "size": r"\b(size guide|sizing|dimensions|measurements|width|height|length|capacity)\b",
        "material": r"\b(material|finish|fabric|care|specifications|specs)\b",
        "financing": r"\b(financ|installment|instalment|bnpl|monthly payments?)\b",
        "fitment": r"\b(vehicle|fitment|compatib|make|model)\b",
    }
    positive = None
    for key, pattern in terms.items():
        if key in observation:
            positive = bool(re.search(pattern, text, re.I))
            break
    if positive is None:
        return None, "No deterministic DOM rule exists for this observation; review required."
    scope = next((value for key, value in terms.items() if key in observation), "the rendered page")
    if positive:
        return False, f"Rendered DOM contains evidence matching {scope!r}; the checklist problem was not observed."
    return True, f"Rendered DOM contained no visible {scope!r} evidence across the captured pages."


def answer_dom_rows(rows: list[dict], snapshots: dict) -> dict[int, dict]:
    result = {}
    for row in rows:
        exists, evidence = _dom_presence(row, snapshots)
        result[row["#"]] = {
            "exists": "unsure" if exists is None else ("Y" if exists else "N"),
            "confidence": "low" if exists is None else ("high" if exists is False else "medium"),
            "evidence": evidence,
        }
    return result


def review_result(row: dict, output_folder: Path) -> dict:
    output_folder = Path(output_folder)
    desktop = output_folder / "desktop-home.png"
    mobile = output_folder / "mobile-home.png"
    return {
        "exists": "review",
        "confidence": "human review",
        "evidence": f"Human review required; desktop screenshot: {desktop.as_posix()}; mobile screenshot: {mobile.as_posix()}.",
    }


def _api_payload(url: str, key: str, strategy: str = "mobile") -> dict:
    query = urllib.parse.urlencode({"url": url, "key": key, "strategy": strategy})
    request = urllib.request.Request("https://www.googleapis.com/pagespeedonline/v5/runPagespeed?" + query)
    with urllib.request.urlopen(request, timeout=120) as response:
        return {"http_status": response.status, "response": json.load(response)}


def run_pagespeed(url: str, key: str) -> dict[int, dict]:
    try:
        payload = None
        last_error = None
        for _ in range(3):
            try:
                payload = _api_payload(url, key)
                break
            except Exception as exc:
                last_error = exc
        if payload is None:
            raise last_error
    except Exception as exc:
        evidence = f"PageSpeed request failed ({type(exc).__name__}); API evidence unavailable."
        return {number: {"exists": "unsure", "confidence": "low", "evidence": evidence} for number in (56, 57, 58, 59)}
    results = {}
    for number in (56, 57, 58, 59):
        results[number] = evaluate_pagespeed(number, payload)
    return results


def _applicable(row: dict, niche: str) -> bool:
    applies = str(row.get("Applies to") or "")
    return "All stores" in applies or niche.lower() in applies.lower()


def assemble(rows: list[dict], store: dict, snapshots: dict, output_folder: Path, api_results: dict[int, dict]) -> list[dict]:
    dom_rows = [row for row in rows if row["Method"] == "DOM"]
    dom_results = answer_dom_rows(dom_rows, snapshots)
    assembled = []
    for row in rows:
        if row["Method"] in ("Vision", "Manual"):
            result = review_result(row, output_folder)
        elif not _applicable(row, str(store.get("Niche") or "")):
            result = {"exists": "NA", "confidence": "high", "evidence": "Not applicable to the store niche listed in the input workbook."}
        elif row["Method"] == "DOM":
            result = dom_results[row["#"]]
        elif row["Method"] == "API":
            result = api_results[row["#"]]
        else:
            result = {"exists": "unsure", "confidence": "low", "evidence": "Unsupported method."}
        assembled.append({**row, "Exists?": result["exists"], "Confidence": result["confidence"],
                          "Evidence": result["evidence"],
                          "Email-ready line": row["Email-ready line"] if result["exists"] == "Y" else None})
    return assembled


def write_report(path: Path, store: dict, rows: list[dict]) -> None:
    workbook = load_workbook(path) if path.exists() else __import__("openpyxl").Workbook()
    if workbook.sheetnames == ["Sheet"] and workbook["Sheet"].max_row == 1:
        del workbook["Sheet"]
    title = str(store.get("Store Name") or "Audited Store")[:31]
    if title in workbook.sheetnames:
        del workbook[title]
    ws = workbook.create_sheet(title)
    headers = ["#", "Category", "Observation", "How to check", "Email-ready line", "Impact", "Applies to", "Method", "Exists?", "Confidence", "Evidence"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = min(80, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)


def update_input(path: Path, store_url: str, output_link: str) -> None:
    workbook = load_workbook(path)
    ws = workbook["Sites"]
    header_cells = next(row for row in ws.iter_rows() if row[1].value == "Store URL")
    header_row = header_cells[0].row
    headers = {cell.value: cell.column for cell in header_cells}
    for row in range(header_row + 1, ws.max_row + 1):
        if ws.cell(row, headers["Store URL"]).value == store_url and ws.cell(row, headers["Status"]).value == "Pending":
            ws.cell(row, headers["Status"]).value = "Done"
            ws.cell(row, headers["Date Audited"]).value = datetime.now(timezone.utc).date().isoformat()
            ws.cell(row, headers["Output File / Link"]).value = output_link
            workbook.save(path)
            return
    raise ValueError("Pending store row disappeared before update")


def audit_store(project: Path) -> Path:
    load_dotenv(project / ".env")
    checklist = load_checklist(project / "cro-audit-checklist.xlsx")
    store = first_pending_store(project / "sites-to-audit.xlsx")
    output_folder = project / "output" / slugify(str(store["Store Name"]))
    output_folder.mkdir(parents=True, exist_ok=True)
    evidence = browser.collect(str(store["Store URL"]), output_folder)
    if evidence.get("status") == "could not audit":
        raise RuntimeError(evidence.get("status"))
    api_results = run_pagespeed(str(store["Store URL"]), os.getenv("PAGESPEED_API_KEY", ""))
    report = project / "output" / "cro-audit-results.xlsx"
    rows = assemble(checklist, store, evidence.get("snapshots", {}), output_folder, api_results)
    write_report(report, store, rows)
    update_input(project / "sites-to-audit.xlsx", str(store["Store URL"]), report.as_posix())
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--smoke", action="store_true")
    args = parser.parse_args()
    if args.smoke:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser_instance = p.chromium.launch()
            page = browser_instance.new_page()
            page.set_content("<title>scaffold ready</title>")
            print(page.title())
            browser_instance.close()
        return
    print(audit_store(args.project))


if __name__ == "__main__":
    main()
