"""Independent artifact contract check. Run with the project's Python environment."""
import json
import re
import struct
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

OUTPUT = Path(__file__).resolve().parent
PROJECT = OUTPUT.parent
SPEC = load_workbook(PROJECT / 'cro-audit-checklist.xlsx', data_only=True)
REPORT = OUTPUT / 'cro-audit-results.xlsx'


def verify():
    assert REPORT.is_file(), f'Missing output workbook: {REPORT}'
    report = load_workbook(REPORT, data_only=True)
    spec_rows = list(SPEC['Master Checklist'].values)
    source = {row[0]: dict(zip(spec_rows[0], row)) for row in spec_rows[1:]}
    candidates = []
    for sheet in report:
        rows = list(sheet.values)
        if rows and rows[0][0] == '#' and 'Evidence' in rows[0]:
            candidates.append((sheet, rows))
    assert len(candidates) == 1, 'Phase 1 must contain exactly one audited store sheet'
    sheet, rows = candidates[0]
    result = [dict(zip(rows[0], row)) for row in rows[1:] if row[0] is not None]
    assert len(result) == 92, f'Expected 92 rows, got {len(result)}'
    assert [r['#'] for r in result] == list(range(1, 93)), 'Missing, duplicated or reordered checklist IDs'
    status_header = next(c for c in rows[0] if c.startswith('Exists?'))
    assert 'Confidence' in rows[0]
    for row in result:
        original = source[row['#']]
        for field in ('Category', 'Observation', 'How to check', 'Impact', 'Applies to', 'Method'):
            assert row[field] == original[field], f'Changed specification at #{row["#"]} / {field}'
        state = row[status_header]
        assert state in ('Y', 'N', 'NA', 'unsure', 'review'), (row['#'], state)
        assert row['Confidence'] is not None, f'Missing confidence #{row["#"]}'
        assert state == 'NA' or (isinstance(row['Evidence'], str) and row['Evidence'].strip()), f'Missing evidence #{row["#"]}'
        if row['Method'] in ('Vision', 'Manual'):
            assert state == 'review', f'Model/manual row was judged: #{row["#"]}'
            assert '.png' in row['Evidence'], f'Missing screenshot reference #{row["#"]}'
            assert 'desktop' in row['Evidence'].lower() and 'mobile' in row['Evidence'].lower(), f'Missing both devices #{row["#"]}'
        else:
            assert state != 'review', f'Automated method incorrectly routed to review #{row["#"]}'
        if state == 'Y':
            assert row['Email-ready line'] == original['Email-ready line'], f'Email differs #{row["#"]}'
        else:
            assert row['Email-ready line'] in (None, ''), f'Email emitted for non-Y #{row["#"]}'
    counts = Counter(r['Method'] for r in result)
    assert counts == {'DOM':61,'API':4,'Vision':20,'Manual':7}, counts
    assert Counter(r[status_header] for r in result)['review'] == 27
    pngs = list(OUTPUT.rglob('*.png'))
    assert pngs, 'No screenshots'
    for path in pngs:
        with path.open('rb') as file:
            header = file.read(24)
        assert header[:8] == b'\x89PNG\r\n\x1a\n', f'Invalid PNG {path}'
        width, height = struct.unpack('>II', header[16:24])
        assert width > 0 and height > 0, f'Empty screenshot {path}'
    result_by_id = {r['#']: r for r in result}
    expectations = json.loads((OUTPUT / 'lo-and-co-interiors/qa-expectations.json').read_text(encoding='utf-8'))
    comparisons = []
    for check in expectations['checks']:
        artifact = OUTPUT / 'lo-and-co-interiors' / check['evidence_file']
        assert artifact.is_file(), f'Missing independent QA evidence: {artifact}'
        actual = result_by_id[check['id']][status_header]
        comparisons.append({**check, 'actual':actual, 'matches':actual == check['expected']})
    matched = sum(c['matches'] for c in comparisons)
    data = {'structural_contract':'passed','rows':len(result),'routing':dict(counts),
            'statuses':dict(Counter(r[status_header] for r in result)),
            'png_files_valid':len(pngs),'qa_matches':matched,'qa_compared':len(comparisons),
            'qa_comparisons':comparisons}
    (OUTPUT / 'delivery-verification.json').write_text(json.dumps(data, indent=2), encoding='utf-8')
    print(json.dumps({k:v for k,v in data.items() if k != 'qa_comparisons'},indent=2))
    assert matched >= 10, f'Only {matched} independently matched judgments; need at least 10'
    return data


if __name__ == '__main__':
    verify()
